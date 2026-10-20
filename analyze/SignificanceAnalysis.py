from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import beta
from statsmodels.stats.contingency_tables import mcnemar
from statsmodels.stats.multitest import multipletests
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
INPUT_XLSX = PROJECT_ROOT / 'classifyOutcome' / 'summarizedData.xlsx'
OUTPUT_DIR = PROJECT_ROOT / 'analyzeOutcome' / 'Statistical_Significance'

N_EXPECTED = 675
N_BOOTSTRAP = 10000
SEED = 42
ALPHA = 0.05
EXACT_THRESHOLD = 25

SHEETS = {
    'BERT': 'BERT',
    'Confidence Cascade Routing': 'Cascade',
    'Single-Model One-Shot (doubao)': 'Doubao One-shot',
    'Hierarchical Binary Routing': 'Hierarchical',
    'Annotation-Enhanced': 'Annotation-enhanced',
}
PRIMARY = ['BERT', 'Doubao One-shot', 'Hierarchical', 'Cascade']
ID_COLUMNS = ['sample_id', 'sample id', 'sample_index', 'sample index', 'original_sample_id']
TEXT_COLUMNS = ['requirement', 'requirement_text', 'requirement_content', 'text']
LABEL_COLUMNS = ['true_label', 'human_label', 'gold_label', 'reference_label', 'ground_truth']
CORRECT_COLUMNS = ['is_correct', 'iscorrect', 'bert_correct', 'is correct']
PRED_COLUMNS = ['model_prediction', 'predicted_label', 'final_predicted_label', 'bert_pred', 'prediction']


def normalize_columns(df):
    df = df.copy()
    df.columns = [str(x).strip().lower().replace('-', '_') for x in df.columns]
    return df


def find_column(df, options):
    return next((c for c in options if c in df.columns), None)


def normalized_id(value):
    if pd.isna(value):
        raise ValueError('Missing sample_id')
    if isinstance(value, (int, np.integer)):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = str(value).strip()
    if not result:
        raise ValueError('Empty sample_id')
    return result


def normalize_label(value):
    if pd.isna(value):
        return None
    return str(value).strip().lower()


def normalize_text(value):
    if pd.isna(value):
        return None
    return ' '.join(str(value).split())


def binary_correct(value):
    if pd.isna(value):
        raise ValueError('Missing correctness value')
    if isinstance(value, (bool, np.bool_)):
        return int(value)
    if isinstance(value, (int, np.integer)) and value in (0, 1):
        return int(value)
    if isinstance(value, (float, np.floating)) and value in (0.0, 1.0):
        return int(value)
    text = str(value).strip().lower()
    if text in ('1', 'true', 'yes', 'y', 'correct'):
        return 1
    if text in ('0', 'false', 'no', 'n', 'incorrect', 'wrong'):
        return 0
    raise ValueError(f'Invalid correctness value: {value!r}')


def load_sheet(sheet_name, short_name):
    df = normalize_columns(pd.read_excel(INPUT_XLSX, sheet_name=sheet_name))
    id_col = find_column(df, ID_COLUMNS)
    if id_col is None:
        raise ValueError(f'{sheet_name}: No stable sample_id. Fix summarize_all.py; row-order pairing is forbidden.')
    df['sample_id'] = df[id_col].map(normalized_id)
    if df['sample_id'].duplicated().any():
        duplicates = df.loc[df['sample_id'].duplicated(), 'sample_id'].head().tolist()
        raise ValueError(f'{sheet_name}: Duplicate sample IDs: {duplicates}')
    ccol = find_column(df, CORRECT_COLUMNS)
    if ccol is None:
        raise ValueError(f'{sheet_name}: Missing binary correctness column; found {list(df.columns)}')
    try:
        df['correct'] = df[ccol].map(binary_correct).astype(np.int8)
    except ValueError as exc:
        raise ValueError(f'{sheet_name}: {exc}') from exc
    text_col = find_column(df, TEXT_COLUMNS)
    label_col = find_column(df, LABEL_COLUMNS)
    pred_col = find_column(df, PRED_COLUMNS)
    df['text_check'] = df[text_col].map(normalize_text) if text_col else None
    df['label_check'] = df[label_col].map(normalize_label) if label_col else None
    if label_col and pred_col:
        derived = (df[label_col].map(normalize_label) == df[pred_col].map(normalize_label)).astype(int)
        mismatch = derived != df['correct']
        if mismatch.any():
            raise ValueError(f'{sheet_name}: {int(mismatch.sum())} correctness flags disagree with gold/prediction labels')
    return df[['sample_id', 'correct', 'text_check', 'label_check']].set_index('sample_id').rename(columns={'correct': short_name})


def align_and_audit():
    frames = {short: load_sheet(sheet, short) for sheet, short in SHEETS.items()}
    ref = frames['BERT']
    ref_ids = set(ref.index)
    if len(ref_ids) != N_EXPECTED:
        raise ValueError(f'BERT has {len(ref_ids)} samples, expected {N_EXPECTED}')
    audit = []
    for name, frame in frames.items():
        ids = set(frame.index)
        missing, extra = ref_ids - ids, ids - ref_ids
        if missing or extra:
            raise ValueError(f'{name}: IDs differ from BERT; missing={list(sorted(missing))[:5]}, extra={list(sorted(extra))[:5]}')
        aligned = frame.reindex(ref.index)
        for field in ['text_check', 'label_check']:
            left, right = ref[field], aligned[field]
            shared = left.notna() & right.notna()
            mismatches = int((left[shared] != right[shared]).sum())
            if mismatches:
                raise ValueError(f'{name}: {mismatches} mismatches in {field}; ID mapping is unsafe')
            audit.append({'Method': name, 'Check': field, 'Compared rows': int(shared.sum()), 'Mismatches': mismatches,
                          'Note': 'NOT VERIFIED (field absent)' if not shared.any() else 'PASS'})
    wide = pd.DataFrame({name: frame.reindex(ref.index)[name] for name, frame in frames.items()})
    if wide.isna().any().any():
        raise ValueError('Missing correctness after ID alignment')
    wide = wide.sort_index()
    for name in wide.columns:
        audit.append({'Method': name, 'Check': 'sample_id + binary correctness', 'Compared rows': len(wide),
                      'Mismatches': 0, 'Note': 'PASS'})
    return wide, pd.DataFrame(audit)


def paired_bootstrap(a, b, rng):
    diff = a.astype(float) - b.astype(float)
    n = len(diff)
    draws = np.empty(N_BOOTSTRAP, dtype=float)
    for k in range(N_BOOTSTRAP):
        ix = rng.integers(0, n, n)
        draws[k] = diff[ix].mean()
    low, high = np.quantile(draws, [ALPHA / 2, 1 - ALPHA / 2])
    return 100 * float(diff.mean()), 100 * float(low), 100 * float(high)


def accuracy_bootstrap(a, rng):
    n = len(a)
    draws = np.empty(N_BOOTSTRAP)
    for k in range(N_BOOTSTRAP):
        draws[k] = a[rng.integers(0, n, n)].mean()
    lo, hi = np.quantile(draws, [ALPHA / 2, 1 - ALPHA / 2])
    return 100 * float(lo), 100 * float(hi)


def cohens_h(a, b):
    return 2 * (math.asin(math.sqrt(a)) - math.asin(math.sqrt(b)))


def paired_test(a, b, name_a, name_b, rng):
    n = len(a)
    both = int(((a == 1) & (b == 1)).sum())
    a_only = int(((a == 1) & (b == 0)).sum())
    b_only = int(((a == 0) & (b == 1)).sum())
    neither = int(((a == 0) & (b == 0)).sum())
    assert both + a_only + b_only + neither == n
    assert int(a.sum() - b.sum()) == a_only - b_only
    discordant = a_only + b_only
    exact = discordant < EXACT_THRESHOLD
    if discordant == 0:
        p_value, statistic = 1.0, 0.0
    else:
        res = mcnemar([[both, a_only], [b_only, neither]], exact=exact, correction=not exact)
        p_value, statistic = float(res.pvalue), float(res.statistic)
    delta, lo, hi = paired_bootstrap(a, b, rng)
    return {
        'Method A': name_a, 'Method B': name_b, 'N paired': n,
        'Accuracy A (%)': 100 * float(a.mean()), 'Accuracy B (%)': 100 * float(b.mean()),
        'Delta A-B (pp)': delta, 'Paired CI low (pp)': lo, 'Paired CI high (pp)': hi,
        'Both correct': both, 'A only correct': a_only, 'B only correct': b_only,
        'Both wrong': neither, 'Discordant pairs': discordant,
        'McNemar type': 'exact' if exact else 'chi-square continuity corrected',
        'McNemar statistic': statistic, 'Raw p': p_value,
        'Matched OR (Haldane 0.5)': (a_only + 0.5) / (b_only + 0.5),
        "Cohen's h": cohens_h(float(a.mean()), float(b.mean())),
    }


def save_excel(df, path):
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Results')
        ws = writer.sheets['Results']
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[1].height = 34
        for col in ws.columns:
            letter = col[0].column_letter
            max_width = min(45, max(14, max(len(str(c.value or '')) for c in list(col)[:200]) + 2))
            ws.column_dimensions[letter].width = max_width


def main():
    if not INPUT_XLSX.is_file():
        raise FileNotFoundError(f'Missing input: {INPUT_XLSX}. Run summarize_all.py first.')
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wide, audit = align_and_audit()
    rng = np.random.default_rng(SEED)
    acc_rows = []
    for name in list(PRIMARY) + ['Annotation-enhanced']:
        a = wide[name].to_numpy(dtype=np.int8)
        lo, hi = accuracy_bootstrap(a, rng)
        acc_rows.append({'Method': name, 'N': len(a), 'Correct': int(a.sum()),
                         'Accuracy (%)': 100 * float(a.mean()), 'CI low (%)': lo, 'CI high (%)': hi,
                         'Evaluation': 'Human-label-assisted final agreement' if name == 'Annotation-enhanced' else 'Autonomous'})
    acc_df = pd.DataFrame(acc_rows)
    pairs = [paired_test(wide[x].to_numpy(), wide[y].to_numpy(), x, y, rng) for x, y in combinations(PRIMARY, 2)]
    pair_df = pd.DataFrame(pairs)
    for method in ['holm', 'bonferroni']:
        rejected, corrected, _, _ = multipletests(pair_df['Raw p'].to_numpy(), alpha=ALPHA, method=method)
        pair_df[f'{method.title()} p'] = corrected
        pair_df[f'{method.title()} significant'] = rejected
    pair_df['Significance (Holm)'] = pair_df['Holm p'].map(lambda p: '***' if p < .001 else '**' if p < .01 else '*' if p < .05 else 'n.s.')
    pair_df = pair_df.sort_values('Delta A-B (pp)', key=lambda x: x.abs(), ascending=False).reset_index(drop=True)

    # Annotation is an assisted workflow, NOT included in the six primary significance tests.
    annotation = wide['Annotation-enhanced'].to_numpy(dtype=np.int8)
    assisted_rows = []
    for name in PRIMARY:
        a = wide[name].to_numpy(dtype=np.int8)
        delta, lo, hi = paired_bootstrap(annotation, a, rng)
        assisted_rows.append({'Comparison': f'Annotation-enhanced minus {name}', 'N': len(a),
                              'Agreement difference (pp)': delta, 'Paired CI low (pp)': lo,
                              'Paired CI high (pp)': hi,
                              'Interpretation': 'Descriptive only: reference labels were used for correction'})
    annotation_df = pd.DataFrame(assisted_rows)

    save_excel(acc_df, OUTPUT_DIR / 'Bootstrap_Confidence_Intervals.xlsx')
    save_excel(pair_df, OUTPUT_DIR / 'Pairwise_McNemar_Tests.xlsx')
    save_excel(pair_df[['Method A', 'Method B', 'Delta A-B (pp)', 'Paired CI low (pp)', 'Paired CI high (pp)',
                        'Matched OR (Haldane 0.5)', "Cohen's h"]], OUTPUT_DIR / 'Effect_Sizes.xlsx')
    save_excel(pair_df[['Method A', 'Method B', 'Accuracy A (%)', 'Accuracy B (%)', 'Delta A-B (pp)',
                        'Paired CI low (pp)', 'Paired CI high (pp)', 'Raw p', 'Holm p',
                        'Holm significant', 'Significance (Holm)']], OUTPUT_DIR / 'Summary_Table_With_Significance.xlsx')
    save_excel(annotation_df, OUTPUT_DIR / 'Annotation_Assisted_Descriptive.xlsx')
    save_excel(audit, OUTPUT_DIR / 'Sample_Alignment_Audit.xlsx')

    fig, ax = plt.subplots(figsize=(8, 6))
    matrix = np.zeros((len(PRIMARY), len(PRIMARY)))
    for _, r in pair_df.iterrows():
        i, j = PRIMARY.index(r['Method A']), PRIMARY.index(r['Method B'])
        matrix[i, j] = matrix[j, i] = -np.log10(max(float(r['Holm p']), 1e-300))
    im = ax.imshow(matrix, cmap='viridis')
    ax.set_xticks(range(len(PRIMARY)), PRIMARY, rotation=30, ha='right')
    ax.set_yticks(range(len(PRIMARY)), PRIMARY)
    for i in range(len(PRIMARY)):
        for j in range(len(PRIMARY)):
            ax.text(j, i, '-' if i == j else f'{matrix[i,j]:.1f}', ha='center', va='center',
                    color='white' if matrix[i,j] > matrix.max() / 2 else 'black')
    ax.set_title('Pairwise significance: -log10(Holm-adjusted p)')
    fig.colorbar(im, ax=ax, label='-log10(adjusted p)')
    fig.tight_layout()
    fig.savefig(OUTPUT_DIR / 'Significance_Heatmap.png', dpi=300)
    plt.close(fig)

    provenance = {
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'input_file': str(INPUT_XLSX),
        'input_sha256': hashlib.sha256(INPUT_XLSX.read_bytes()).hexdigest(),
        'n_matched': len(wide), 'n_bootstrap': N_BOOTSTRAP, 'random_seed': SEED,
        'confidence_interval': 'paired percentile bootstrap (same resampled IDs for both methods)',
        'primary_family': PRIMARY, 'n_primary_comparisons': len(pair_df),
        'multiple_comparison_correction': 'Holm (primary), Bonferroni (sensitivity)',
        'mcnemar': f'exact if discordant < {EXACT_THRESHOLD}; otherwise continuity-corrected chi-square',
        'annotation_policy': 'Assisted method excluded from primary autonomous comparisons',
        'alignment_note': 'IDs matched exactly; text and labels compared wherever present; inspect audit file for missing fields',
    }
    (OUTPUT_DIR / 'Analysis_Metadata.json').write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding='utf-8')
    print('\n=== Classification accuracy ===')
    print(acc_df.to_string(index=False, float_format=lambda x: f'{x:.4f}'))
    print('\n=== Primary paired tests ===')
    print(pair_df[['Method A', 'Method B', 'Delta A-B (pp)', 'Paired CI low (pp)',
                   'Paired CI high (pp)', 'Raw p', 'Holm p']].to_string(index=False, float_format=lambda x: f'{x:.6g}'))
    print(f'\nSaved results to: {OUTPUT_DIR}')


if __name__ == '__main__':
    main()
